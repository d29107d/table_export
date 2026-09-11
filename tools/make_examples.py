"""生成 ``example/`` 下的示例表格与导出产物。

用法::

    python tools/make_examples.py

做两件事：

1. 在 ``example/import/`` 里生成示例 xlsx（覆盖全部表类型与字段类型）
2. 用本工具的导出逻辑把它们导到 ``example/client/`` 与 ``example/server/``

示例表就是"格式说明书"——想知道某一行某一列是干什么的，看这个脚本比看文档准。
"""

import os
import sys

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import openpyxl

from core.excel_reader import list_excel_files, load_excel
from core.exporter import export_table

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
IMPORT_DIR = os.path.join(ROOT, "example", "import")
CLIENT_DIR = os.path.join(ROOT, "example", "client")
SERVER_DIR = os.path.join(ROOT, "example", "server")

#: 示例统一用 UTF-8 导出，方便在 GitHub 上直接看。
#: 实际项目里前后端可以各用各的编码（界面里分别选）。
ENCODING = "utf-8"

#: 生成产物时只清理自己产出的文件，不动目录里的其他东西
_OUTPUT_PREFIX = "cfg_example_"


class _Skip:
    """占位：这一格"根本不存在"（不写入单元格），与空字符串区分开。"""


SKIP = _Skip()


# ── 写表 ─────────────────────────────────────────────────────────


def _set_meta(ws, kind, filename, key_count=0, header="return {", footer="}"):
    """写 1~3 行的表级元信息（两种表类型都一样）。"""
    ws["A1"] = "导出类型"
    ws["B1"] = kind
    ws["A2"] = "导出文件"
    ws["B2"] = filename
    ws["A3"] = "key数量"
    ws["B3"] = key_count
    ws["D1"] = "导出文件头"
    ws["E1"] = header
    ws["D2"] = "导出文件尾"
    ws["E2"] = footer


def write_base_sheet(wb, title, filename, key_count, columns, rows,
                     header="return {", footer="}"):
    """写一张 base 表（一行一条数据）。

    ``columns``: ``(注释, scope, 类型, 字段名)`` 列表
    ``rows``: 每行一个列表，元素顺序与 ``columns`` 一致；用 ``SKIP`` 表示空着这一格
    """
    ws = wb.create_sheet(title)
    _set_meta(ws, "base", filename, key_count, header, footer)

    for col, (note, scope, ftype, fname) in enumerate(columns, start=1):
        ws.cell(row=5, column=col, value=note)
        ws.cell(row=6, column=col, value=scope)
        ws.cell(row=7, column=col, value=ftype)
        ws.cell(row=8, column=col, value=fname)

    for r, row in enumerate(rows, start=9):
        for c, value in enumerate(row, start=1):
            if value is SKIP:
                continue
            ws.cell(row=r, column=c, value=value)
    return ws


def write_tiny_sheet(wb, title, filename, fields, header="return {", footer="}"):
    """写一张 tiny 表（一行一个字段，整个文件就一条记录）。

    ``fields``: ``(注释, scope, 类型, 字段名, 值)`` 列表；用 ``SKIP`` 表示整行留空。
    """
    ws = wb.create_sheet(title)
    _set_meta(ws, "tiny", filename, 0, header, footer)

    for col, head in enumerate(("配置备注", "导出参数", "值类型", "字段名", "值"), start=1):
        ws.cell(row=5, column=col, value=head)

    for r, field in enumerate(fields, start=6):
        if field is SKIP:
            continue
        note, scope, ftype, fname, value = field
        ws.cell(row=r, column=1, value=note)
        ws.cell(row=r, column=2, value=scope)
        ws.cell(row=r, column=3, value=ftype)
        ws.cell(row=r, column=4, value=fname)
        if value is not SKIP:
            ws.cell(row=r, column=5, value=value)
    return ws


def _new_workbook():
    """新建工作簿并丢掉默认的 Sheet。"""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    return wb


# ── 各示例文件 ───────────────────────────────────────────────────


def build_types_and_scopes():
    """01：四种字段类型 × 四种 scope。"""
    wb = _new_workbook()

    columns = [
        ("唯一 ID",        "sc", "number", "id"),
        ("名称",           "c",  "string", "name"),
        ("服务端备注",     "s",  "string", "server_note"),
        ("描述",           "sc", "string", "desc"),
        ("权重",           "cs", "number", "weight"),
        ("价格",           "c",  "number", "price"),
        ("图标",           "c",  "string", "icon"),
        ("掉落表",         "s",  "table",  "drops"),
        ("自定义参数",     "sc", "any",    "extra"),
        ("是否上架",       "c",  "any",    "on_sale"),
    ]
    rows = [
        [1, "长剑",  "新手村产出", "一把普通的长剑", 100, 10,  "icon_1001", "{{1001, 2}, {1002, 1}}", "nil",  True],
        [2, "铁盾",  "铁匠铺出售", "结实的铁盾",     80,  50,  "icon_1002", SKIP,                     "100+50", False],
        [3, "药水",  "商店购买",   SKIP,             200, 5,   "icon_1003", "{{2001, 5}}",            "nil",  True],
        [4, "卷轴",  "活动奖励",   "使用后随机传送", 50,  SKIP, "icon_1004", "{{3001, 1}}",           "{quality=3}", True],
    ]
    write_base_sheet(wb, "Item", "cfg_example_item.lua", 1, columns, rows)

    equip_columns = [
        ("装备 ID",  "sc", "number", "id"),
        ("部位",     "sc", "string", "part"),
        ("攻击",     "sc", "number", "attack"),
        ("防御",     "sc", "number", "defense"),
        ("强化上限", "s",  "number", "max_level"),
        ("套装 ID",  "c",  "number", "suit_id"),
    ]
    equip_rows = [
        [100, "weapon", 12, 0,  15, 1],
        [101, "armor",  0,  18, 15, 1],
        [102, "helmet", 3,  8,  10, 2],
    ]
    write_base_sheet(wb, "Equip", "cfg_example_equip.lua", 1, equip_columns, equip_rows)

    wb.save(os.path.join(IMPORT_DIR, "01_types_and_scopes.xlsx"))
    wb.close()


def build_keys_and_layout():
    """02：key 数量 0 / 1 / 2（决定输出是列表还是多级嵌套），外加会被跳过的两种页。"""
    wb = _new_workbook()

    # key 数量 2 -> [shop_id] = { [item_id] = { ... } }
    shop_columns = [
        ("商店 ID",   "sc", "number", "shop_id"),
        ("商品 ID",   "sc", "number", "item_id"),
        ("售价",      "c",  "number", "price"),
        ("库存",      "s",  "number", "stock"),
        ("商品名",    "sc", "string", "name"),
    ]
    shop_rows = [
        [1, 1001, 10, 999, "长剑"],
        [1, 1002, 50, 999, "铁盾"],
        [2, 1001, 12, 100, "长剑"],
        [2, 2001, 5,  200, "药水"],
    ]
    write_base_sheet(wb, "Shop", "cfg_example_shop.lua", 2, shop_columns, shop_rows)

    # key 数量 0 -> 每条数据一个匿名 table 元素
    tips_columns = [
        ("ID",   "c", "number", "id"),
        ("文案", "c", "string", "text"),
    ]
    tips_rows = [
        [1, "欢迎来到示例服务器"],
        [2, "按 H 打开帮助"],
        [3, "每天 5 点重置日常"],
    ]
    write_base_sheet(wb, "Tips", "cfg_example_tips.lua", 0, tips_columns, tips_rows)

    # 会被跳过的两种页（都**不会**产生 lua 文件）
    ws = wb.create_sheet("Readme")
    ws["A1"] = "说明页"
    ws["B1"] = "这部分不是配置"
    ws["A3"] = "B1 既不是 base 也不是 tiny，整个 sheet 会被忽略。"

    empty_columns = [
        ("ID",   "sc", "number", "id"),
        ("名称", "sc", "string", "name"),
    ]
    write_base_sheet(wb, "EmptyTable", "cfg_example_empty.lua", 1, empty_columns, [])

    wb.save(os.path.join(IMPORT_DIR, "02_keys_and_layout.xlsx"))
    wb.close()


def build_tiny_config():
    """03：tiny 表——整个文件只有一条记录，一行一个字段。"""
    wb = _new_workbook()

    fields = [
        ("每日免费复活次数", "sc", "number", "free_revive_count", 3),
        ("复活消耗文本",     "c",  "string", "revive_cost_text", "消耗 %d 元宝复活"),
        ("入口 NPC",         "s",  "table",  "entrance_npc",     '{10000, "复活使者", 1503}'),
        ("功能开关",         "c",  "any",    "enable",           True),
        ("额外配置",         "sc", "any",    "extra",            "nil"),
        ("空串开关",         "c",  "any",    "blank_flag",       " "),
        SKIP,  # 中间留一行空行：tiny 表会跳过空行继续往下读
        ("活动倍率",         "s",  "number", "activity_rate",    1.5e3),
        ("公告",             "c",  "string", "notice",           "欢迎光临"),
    ]
    write_tiny_sheet(wb, "Settings", "cfg_example_setting.lua", fields)

    wb.save(os.path.join(IMPORT_DIR, "03_tiny_config.xlsx"))
    wb.close()


def build_edge_cases():
    """04：各种边界情形。"""
    wb = _new_workbook()

    columns = [
        ("ID",              "sc", "number", "id"),
        ("多行文本",        "c",  "string", "multiline"),
        ("含 ]] 的文本",    "c",  "string", "square_brackets"),
        ("以 ] 结尾的文本",  "c",  "string", "trailing_bracket"),
        ("文本：满/空格/空着", "sc", "string", "optional_text"),
        ("数字：满/空格/空着", "c", "number", "blank_number"),
        ("表：满/空格/空着",  "c", "table",  "blank_table"),
        ("布尔开关",        "sc", "any",    "flag"),
        ("科学计数",        "s",  "number", "scientific"),
        ("负数",            "s",  "number", "negative"),
    ]
    # 第 5/6/7 列刻意排出三种情形：正常值 / 只填了空格（= 空串）/ 整格空着
    # -> 产物分别是 [[ ]]（空串）、nil、{}，而"整格空着"则是该字段根本不输出
    rows = [
        [1, "第一行\n第二行", "a]]b", "[边界]", "有值", 100, "{{1, 2}}",  True,  1.5e3,  -42],
        [2, "单行文本",       "a]b",   "尾]",   " ",    " ",  " ",        False, 2.5e-2, -7],
        [3, SKIP,             "普通",  "正常",  SKIP,  SKIP, SKIP,       "nil",  1e10,   -1],
    ]
    write_base_sheet(wb, "Edge", "cfg_example_edge.lua", 1, columns, rows)

    # 文件头/尾可以自定义：这里包一层 local + return，产物依然是合法 lua
    custom_columns = [
        ("ID",   "sc", "number", "id"),
        ("名称", "sc", "string", "name"),
    ]
    custom_rows = [[1, "甲"], [2, "乙"]]
    write_base_sheet(wb, "CustomHeaderFooter", "cfg_example_custom.lua", 1,
                     custom_columns, custom_rows,
                     header="local cfg = {", footer="}\nreturn cfg")

    wb.save(os.path.join(IMPORT_DIR, "04_edge_cases.xlsx"))
    wb.close()


# ── 导出 ─────────────────────────────────────────────────────────


def clean_outputs():
    """只删自己生成过的产物，别动目录里其他文件。"""
    for d in (CLIENT_DIR, SERVER_DIR):
        os.makedirs(d, exist_ok=True)
        for name in os.listdir(d):
            if name.startswith(_OUTPUT_PREFIX) and name.endswith(".lua"):
                os.remove(os.path.join(d, name))


def export_examples():
    exported, failed = 0, 0
    for path in sorted(list_excel_files(IMPORT_DIR)):
        for info in load_excel(path):
            result = export_table(path, info, CLIENT_DIR, SERVER_DIR, ENCODING, ENCODING)
            for msg in result.success:
                print("  " + msg)
                exported += 1
            for msg in result.failed:
                print("  FAILED " + msg)
                failed += 1
    return exported, failed


def main():
    os.makedirs(IMPORT_DIR, exist_ok=True)
    print("generating example workbooks -> %s" % IMPORT_DIR)
    build_types_and_scopes()
    build_keys_and_layout()
    build_tiny_config()
    build_edge_cases()

    print("exporting -> %s , %s" % (CLIENT_DIR, SERVER_DIR))
    clean_outputs()
    exported, failed = export_examples()

    print("done: %d file(s) written, %d failed" % (exported, failed))
    return 1 if failed else 0


if __name__ == "__main__":
    sys.exit(main())
